"""Read edits made in Word back into the template and the data.

    python scripts/sync_programmes.py                  # every edited programme
    python scripts/sync_programmes.py ai-ml-aws-130h

For each res/programmes/*.docx changed since its last build:
  1. text  -> written back to programmes/<slug>.json (diff printed)
  2. look  -> printed: styles changed in Word, one-off formatting, header/footer,
              page setup. These go into STYLES in build_programmes.py by hand,
              so every programme gets them.
Then rebuild with build_programmes.py --force (the build refuses to overwrite
an edited file until then, so nothing typed in Word is lost).
"""
import difflib, json, re, subprocess, sys, tempfile
from pathlib import Path
from collections import Counter

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml import etree

sys.path.insert(0, str(__import__('pathlib').Path(__file__).parent))
from build_programmes import CHAR_STYLES, LABELS, LABELS_FR, OUT, SRC, STYLES, build_doc, load_built, sha

NOISE = re.compile(r'\s(xmlns(:\w+)?|(w|w14|w15):(rsid\w*|paraId|textId)|wp14:(anchorId|editId))="[^"]*"'
                   r'|\s(id|name)="(\d+|Picture \d+)"')  # ids Word renumbers on every save


def clean(el):
    return NOISE.sub('', etree.tostring(el, encoding='unicode'))


# ---------- 1. text ----------

def parts(p):
    """(label, rest) where label = text in 'Prog Label' runs."""
    lab, rest = [], []
    for r in p.runs:
        (lab if r.style is not None and r.style.name in ('Prog Label', 'Prog Cover Label') else rest).append(r.text)
    return ''.join(lab).strip(), ''.join(rest).strip()


def cell_text(cell, maths_label=None):
    p = cell.paragraphs[0]
    lab, rest = parts(p)
    text = ' '.join(x.text for x in cell.paragraphs if x.text.strip()) if not lab else rest
    return (f'MATH: {rest}', lab) if lab else (text.strip(), None)


def hours(s):
    m = re.search(r'\d+(?:[.,]\d+)?', s or '')
    if not m: return None
    h = float(m.group().replace(',', '.'))
    return int(h) if h == int(h) else h


def read_doc(path, old):
    doc = Document(path)
    g = {k: old[k] for k in ('slug', 'file', 'lang') if k in old}
    labels, warn = {}, []
    g.update(facts=[], phases=[], notes=[])
    where, ph, box = 'cover', None, None
    h2 = 0
    glance_focus = []

    for el in doc.element.body.iterchildren():
        if el.tag == qn('w:tbl'):
            t = Table(el, doc)
            rows = [[c for c in r.cells] for r in t.rows]
            head = [c.text.strip() for c in rows[0]]
            if where == 'cover':
                pass  # author block (logo, name, site): fixed wording, position checked in the look report
            elif where == 'glance':
                labels['glance_cols'] = head
                glance_focus = [r[2].text.strip() for r in rows[1:] if r[0].text.strip()]
                total = [r for r in rows[1:] if not r[0].text.strip()]
                if len(total) > 1:  # buffer row, then total row
                    labels['buffer'] = total[0][1].text.strip(); labels['buffer_focus'] = total[0][2].text.strip()
                    g['buffer'] = hours(total[0][3].text)
                if total: labels['total'] = total[-1][1].text.strip()
            elif where == 'facts':
                g['facts'] = [{'label': r[0].text.strip(), 'value': r[1].text.strip()} for r in rows]
            elif where == 'phase':
                labels['chapter_cols'] = head
                for r in rows[1:]:
                    content, lab = cell_text(r[2])
                    if lab: labels['maths'] = lab
                    ph['chapters'].append({'id': r[0].text.strip(), 'title': r[1].text.strip(), 'content': content,
                                           'format': r[3].text.strip(), 'hours': hours(r[4].text)})
            elif where == 'portfolio':
                labels['portfolio_cols'] = head
                g['portfolio']['rows'] = [{'name': r[1].text.strip(), 'shows': r[2].text.strip(),
                                           'phase': r[3].text.strip(), 'build': r[4].text.strip()} for r in rows[1:]]
            else:
                warn.append(f'table in an unexpected place ({where}): {head}')
            continue
        if el.tag != qn('w:p'):
            continue
        p = Paragraph(el, doc)
        st, text = p.style.name, p.text.strip()
        if not text:
            continue
        lab, rest = parts(p)
        if st == 'Prog Kicker':
            m = re.match(r'(.*?)\s+(\d+)\s*\t\s*(.*)$', p.text.strip())
            if m:
                labels['phase'] = m.group(1)
                ph = {'name': '', 'hours': hours(m.group(3)), 'focus': '', 'intro': '',
                      'chapters': [], 'math': None, 'project': None}
                g['phases'].append(ph); where = 'phase'
            elif where == 'cover':
                labels['kicker'] = text
            else:
                labels['portfolio_kicker'] = text
                g['portfolio'] = {'intro': '', 'rows': []}; where = 'portfolio'
        elif st == 'Prog Cover Kicker': labels['kicker'] = text
        elif st == 'Prog Cover Title': g['title'] = text.replace(' ', ' ')
        elif st == 'Prog Cover Sub': g['subtitle'] = text
        elif st == 'Prog Cover Stat':  # big hours figure; kept in data only if it differs from the phases' sum
            unit = ''.join(r.text for r in p.runs if r.style is not None and r.style.name == 'Prog Cover Unit').strip()
            if unit: labels['hours'] = unit
            g['_hours'] = hours(text.replace(unit, ''))
        elif st == 'Prog Cover Line':
            if 'audience_line' not in g:
                labels['for'] = lab; g['audience_line'] = rest
            else:
                labels['certs'] = lab; g['certifications'] = [c.strip() for c in rest.split('·') if c.strip()]
        elif st == 'Heading 2':
            if where in ('cover', 'overview', 'glance', 'facts'):
                key = ['overview', 'glance', 'facts'][min(h2, 2)]; h2 += 1
                labels[key] = text; where = key
            else:
                labels['notes'] = text; where = 'notes'
        elif st == 'Heading 1':
            if where == 'phase': ph['name'] = text
            elif where == 'portfolio': labels['portfolio'] = text
        elif st == 'Prog Body' and where == 'overview': g['overview'] = (g.get('overview', '') + ' ' + text).strip()
        elif st == 'Prog Intro' and where == 'phase': ph['intro'] = (ph['intro'] + ' ' + text).strip()
        elif st == 'Prog Intro' and where == 'portfolio': g['portfolio']['intro'] = text
        elif st == 'Prog Note' and where == 'phase':
            labels['prereq'] = lab; ph['note'] = rest
        elif st == 'Prog Box Tag':
            m = re.match(r'(.*?)\s+\d+$', text)
            if m or (ph is not None and where == 'phase' and not ph['project'] and text.lower() in
                     (LABELS['final_project'].lower(), LABELS_FR['final_project'].lower())):
                if m: labels['project'] = m.group(1)
                box = ph['project'] = {'title': '', 'brief': '', 'stack': '', 'deliverable': ''}
            else:
                labels['maths_box'] = text; box = 'math'
        elif st == 'Prog Box Title' and isinstance(box, dict): box['title'] = text
        elif st == 'Prog Box Text':
            if box == 'math': ph['math'] = (ph['math'] + ' ' + text) if ph['math'] else text
            elif isinstance(box, dict):
                if not lab: box['brief'] = (box['brief'] + ' ' + text).strip()
                else:  # labelled lines in order: stack, deliverable
                    key = 'stack' if box.pop('_n', 0) == 0 else 'deliverable'
                    labels[key] = lab; box[key] = rest
                    if key == 'stack': box['_n'] = 1
        elif st == 'Prog Bullet' and where == 'notes': g['notes'].append(text)
        elif st == 'Prog Body' and not where == 'overview':
            warn.append(f'loose text (style {st}, after {where}): {text[:60]}')
        else:
            warn.append(f'unknown style "{st}" (after {where}): {text[:60]}')

        if box and st not in ('Prog Box Tag', 'Prog Box Title', 'Prog Box Text'):
            box = None

    shown_hours = g.pop('_hours', None)
    if shown_hours is not None and shown_hours != sum(ph_['hours'] or 0 for ph_ in g['phases']) + (g.get('buffer') or 0):
        g['hours'] = shown_hours
    for ph_, f in zip(g['phases'], glance_focus):
        ph_['focus'] = f
    for ph_ in g['phases']:
        if ph_['project']: ph_['project'].pop('_n', None)
    if len(glance_focus) != len(g['phases']):
        warn.append(f'At a glance has {len(glance_focus)} rows but there are {len(g["phases"])} phases')
    # keep only labels that differ from the template's
    base = LABELS_FR if old.get('lang') == 'fr' else LABELS
    lab = {k: v for k, v in labels.items() if v and base.get(k) != v
           and not (isinstance(v, str) and base.get(k, '').lower() == v.lower())}
    if lab: g['labels'] = {**old.get('labels', {}), **lab}
    elif 'labels' in old: g['labels'] = old['labels']
    order = list(old.keys()) + [k for k in g if k not in old]
    return {k: g[k] for k in order if k in g}, warn


# ---------- 2. look ----------

# styles the template uses (ids): only these, plus brand-new ones, are worth reporting
OURS = {re.sub(r'\W', '', n) for n in list(STYLES) + list(CHAR_STYLES)} | {
    'ProgTable', 'Header', 'Footer', 'ListBullet', 'TableNormal', 'DefaultParagraphFont'}


def style_map(doc):
    return {s.get(qn('w:styleId')): clean(s) for s in doc.styles.element.findall(qn('w:style'))}


def direct(doc):
    """Formatting applied by hand: paragraph/run properties beyond the style."""
    out = Counter(); sample = {}
    for p in doc.element.body.iter(qn('w:p')):
        pPr = p.find(qn('w:pPr'))
        style = ''
        if pPr is not None:
            s = pPr.find(qn('w:pStyle')); style = s.get(qn('w:val')) if s is not None else ''
            extra = [clean(c) for c in pPr if c.tag not in (qn('w:pStyle'), qn('w:rPr'))]
            if extra:
                k = (style, 'paragraph', ' '.join(extra)); out[k] += 1
                sample.setdefault(k, ''.join(t.text or '' for t in p.iter(qn('w:t')))[:50])
        for r in p.findall(qn('w:r')):
            rPr = r.find(qn('w:rPr'))
            if rPr is None: continue
            extra = [clean(c) for c in rPr if c.tag not in (qn('w:rStyle'), qn('w:lang'), qn('w:noProof'))]
            if extra:
                k = (style, 'text', ' '.join(extra)); out[k] += 1
                sample.setdefault(k, ''.join(t.text or '' for t in r.iter(qn('w:t')))[:50])
    return out, sample


EMU_CM = 360000


def placed(doc):
    """Position and size of every floating picture (by name) and floating table, in cm."""
    out = {}
    W = '{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}'
    parts = [doc.element.body] + [x._element for s in doc.sections for x in (s.first_page_header, s.header)]
    for part in parts:
        for a in part.iter(W + 'anchor'):
            name = a.find(W + 'docPr').get('name')
            pos = [a.find(f'{W}position{d}/{W}posOffset') for d in 'HV']
            ext = a.find(W + 'extent')
            crop = a.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}srcRect')
            out[f'picture "{name}"'] = ('x %.2f y %.2f w %.2f h %.2f' % tuple(
                int(v) / EMU_CM for v in (pos[0].text if pos[0] is not None else 0, pos[1].text if pos[1] is not None else 0,
                                          ext.get('cx'), ext.get('cy')))
                + (f' crop {dict(crop.attrib)}' if crop is not None and crop.attrib else ''))
    for i, t in enumerate(doc.element.body.iter(qn('w:tblpPr'))):
        out[f'floating table {i + 1}'] = ' '.join(f'{k.split("}")[1]}={v}' for k, v in sorted(t.attrib.items()))
    return out


def shown(part):
    """A header/footer as it shows: per paragraph, its properties and runs merged by formatting."""
    if part is None: return ''
    out = []
    for p in part.iter(qn('w:p')):
        pPr = p.find(qn('w:pPr'))
        runs = []
        for r in p.iter(qn('w:r')):
            rPr = r.find(qn('w:rPr'))
            fmt = clean(rPr) if rPr is not None else ''
            fmt = re.sub(r'<w:(noProof|lang)[^>]*/>', '', fmt)
            text = ''.join(t.text or '' for t in r.iter(qn('w:t')))
            other = ''.join(clean(c) for c in r if c.tag not in (qn('w:rPr'), qn('w:t')))
            if runs and runs[-1][0] == fmt and not other and not runs[-1][2]:
                runs[-1][1] += text
            else:
                runs.append([fmt, text, other])
        out.append((clean(pPr) if pPr is not None else '') + ''.join(f'<run {f}>{t}{o}</run>' for f, t, o in runs))
    return '\n'.join(out)


def furniture(doc):
    s = doc.sections[0]
    sect = re.sub(r'<w:(header|footer)Reference[^>]*/>', '', clean(s._sectPr))  # ids change on every save
    return {'page setup': sect, 'header': shown(s.header._element), 'footer': shown(s.footer._element),
            'first-page header': shown(s.first_page_header._element)}


def pretty(x):
    return x.replace('><', '>\n<').splitlines()  # one tag per line, enough to diff


def look_report(edited, fresh):
    lines = []
    a, b = style_map(fresh), style_map(edited)
    for sid in sorted(set(a) | set(b)):
        if a.get(sid) == b.get(sid): continue
        if sid not in b: continue  # Word drops unused latent styles; ignore removals
        if sid in a and sid not in OURS: continue  # Word rewrites built-ins the template never uses
        if sid not in a:
            lines.append(f'+ new style {sid}'); lines += ['    ' + l for l in pretty(b[sid])]; continue
        lines.append(f'~ style {sid} changed:')
        lines += ['    ' + l for l in difflib.unified_diff(pretty(a[sid]), pretty(b[sid]), lineterm='', n=0)
                  if not l.startswith(('---', '+++', '@@'))]
    da, _ = direct(fresh); db, sample = direct(edited)
    for k, n in (db - da).items():
        lines.append(f'! hand formatting on {k[1]} ({k[0] or "no style"}) x{n}: {k[2][:160]}  e.g. "{sample[k]}"')
    # objects moved or resized by hand: floating pictures and floating tables
    pa, pb = placed(fresh), placed(edited)
    for k in sorted(set(pa) | set(pb)):
        if pa.get(k) != pb.get(k):
            lines.append(f'~ {k} moved/resized: {pa.get(k, "absent")}  ->  {pb.get(k, "absent")}')
    fa, fb = furniture(fresh), furniture(edited)
    for k in fa:
        if fa[k] != fb[k]:
            lines.append(f'~ {k} changed:')
            lines += ['    ' + l for l in difflib.unified_diff(fa[k].splitlines(), fb[k].splitlines(), lineterm='', n=0)
                      if not l.startswith(('---', '+++', '@@'))][:40]
    return lines


def word_resave(path):
    """Open and save a file in Microsoft Word (Windows), as Word does on any save."""
    ps = (f"$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
          f"$d = $w.Documents.Open('{path}'); $d.Saved = $false; $d.Save(); $d.Close(); $w.Quit()")
    subprocess.run(['powershell', '-NoProfile', '-Command', ps], check=True)


def sync(json_path):
    old = json.loads(json_path.read_text(encoding='utf-8'))
    docx_path = OUT / old['file']
    built = load_built()
    if built.get(old['file']) == sha(docx_path):
        print(f'{old["file"]}: no edits since last build'); return
    print(f'=== {old["file"]}')
    new, warn = read_doc(docx_path, old)
    a = json.dumps(old, ensure_ascii=False, indent=1).splitlines()
    b = json.dumps(new, ensure_ascii=False, indent=1).splitlines()
    diff = list(difflib.unified_diff(a, b, 'before', 'after', lineterm='', n=1))
    print('\n-- text --'); print('\n'.join(diff) if diff else '(no text changes)')
    for w in warn: print('WARNING', w)
    if diff:
        json_path.write_text(json.dumps(new, ensure_ascii=False, indent=1) + '\n', encoding='utf-8')

    with tempfile.TemporaryDirectory() as tmp:
        base = Path(tmp) / 'baseline.docx'
        build_doc(old).save(base)
        word_resave(base)  # so both sides carry Word's own tidying, and only real edits differ
        look = look_report(Document(docx_path), Document(base))
    print('\n-- look --'); print('\n'.join(look) if look else '(no formatting changes)')
    print('\nWhen the look changes are in STYLES: build_programmes.py --force')


if __name__ == '__main__':
    names = sys.argv[1:]
    for f in ([SRC / f'{n}.json' for n in names] if names else sorted(f for f in SRC.glob('*.json') if not f.name.startswith(('.', '_')))):
        sync(f)
