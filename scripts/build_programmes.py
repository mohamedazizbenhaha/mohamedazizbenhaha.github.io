"""Build the training-programme Word files from programmes/*.json.

    python scripts/build_programmes.py                  # every programme
    python scripts/build_programmes.py ai-ml-aws-130h   # one
    python scripts/build_programmes.py --force          # overwrite even if edited in Word

Output: res/programmes/<file>. One template for all programmes. Every block
uses a named Word style (STYLES below), so a style changed in Word can be
carried back here; scripts/sync_programmes.py reads edits back.
A file edited in Word since its last build is never overwritten (see .built.json)
unless --force. Needs python-docx, Pillow and PyMuPDF.
"""
import hashlib, io, json, re, sys
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor, Twips

ROOT = Path(__file__).resolve().parent.parent
SRC, OUT = ROOT / 'programmes', ROOT / 'res' / 'programmes'
BUILT = SRC / '.built.json'
LOGO = ROOT / 'res' / 'my_logo_w.png'
NAME, SITE = 'Mohamed Aziz BEN HAHA', 'mohamedazizbenhaha.netlify.app'

# White page, gold accents. GOLD is the site's gold (rules, fills, big numbers);
# GOLD_TXT is a darker gold that stays readable as small text on white.
INK, MUTE, GOLD, GOLD_TXT = '1A1814', '6B665C', 'D4A024', '9A7413'
PALE, LINE = 'FBF4E2', 'E6DECB'
NIGHT, CREAM, CREAM_MUTE = '0B0B0E', 'F3EFE6', 'A29D92'  # cover page
FONT = 'Calibri'
WIDTH = 17.0  # cm of text width on A4 with 2 cm margins

# Fixed wording of the template; a programme can override any of it in "labels".
LABELS = {
    'kicker': 'Training programme', 'hours': 'Hours', 'for': 'For', 'certs': 'Prepares for',
    'overview': 'Overview', 'glance': 'At a glance', 'facts': 'Key facts',
    'glance_cols': ['#', 'Phase', 'Focus', 'Hours'], 'total': 'Total',
    'phase': 'Phase', 'prereq': 'Prerequisite',
    'chapter_cols': ['#', 'Chapter', 'Content', 'Format', 'Hours'],
    'maths': 'Maths', 'maths_box': 'Maths covered in this phase',
    'project': 'Project', 'stack': 'Stack', 'deliverable': 'Deliverable',
    'portfolio_kicker': 'What participants leave with', 'portfolio': 'Project portfolio',
    'portfolio_cols': ['#', 'Project', 'What it demonstrates', 'Phase', 'Build time'],
    'notes': 'Notes', 'buffer': 'Buffer', 'buffer_focus': 'Catch-up, revision, extra lab time',
    'final_project': 'Final project', 'page': 'Page', 'of': 'of',
}
LABELS_FR = {  # "lang": "fr" in a programme selects these
    'kicker': 'Programme de formation', 'hours': 'Heures', 'for': 'Pour', 'certs': 'Prépare à',
    'overview': 'Présentation', 'glance': 'Vue d’ensemble', 'facts': 'Informations clés',
    'glance_cols': ['#', 'Phase', 'Thèmes', 'Heures'], 'total': 'Total',
    'phase': 'Phase', 'prereq': 'Prérequis',
    'chapter_cols': ['#', 'Chapitre', 'Contenu', 'Format', 'Heures'],
    'maths': 'Maths', 'maths_box': 'Mathématiques abordées dans cette phase',
    'project': 'Projet', 'stack': 'Outils', 'deliverable': 'Livrable',
    'portfolio_kicker': 'Ce que les participants emportent', 'portfolio': 'Portfolio de projets',
    'portfolio_cols': ['#', 'Projet', 'Ce qu’il démontre', 'Phase', 'Durée'],
    'notes': 'Notes', 'buffer': 'Réserve', 'buffer_focus': 'Rattrapage, révision, TP supplémentaires',
    'final_project': 'Projet final', 'page': 'Page', 'of': 'sur',
}

# ---------- styles: the whole look lives here ----------
# size pt · bold · color · caps · track (letter spacing pt) · before/after pt ·
# line (multiple) · keep (with next) · align · tab (right tab at text edge) ·
# border {side: (colour, eighths of a pt, gap pt)} · fill
_BOX = dict(border={'left': (GOLD, 24, 10), 'top': (PALE, 4, 6), 'bottom': (PALE, 4, 6), 'right': (PALE, 4, 10)},
            fill=PALE, line=1.15, indent=(13, 10.5))  # indents = border gap + rule, so the box lines up with the tables
STYLES = {
    'Normal':            dict(size=10.5, color=INK),
    'Heading 1':         dict(size=20, bold=True, color=INK, after=8, keep=True, border={'bottom': (GOLD, 12, 6)}),
    'Heading 2':         dict(size=17, bold=True, color=INK, before=18, after=6, keep=True),
    'Prog Kicker':       dict(size=8.5, bold=True, color=GOLD_TXT, caps=True, track=1.2, after=2, keep=True, tab=True, new_page=True),
    'Prog Cover Kicker': dict(size=9, bold=True, color=GOLD, caps=True, track=2, before=245, after=8),
    'Prog Cover Title':  dict(size=36, bold=True, color=CREAM, after=8, line=0.95),
    'Prog Cover Sub':    dict(size=13, color=CREAM_MUTE, after=22, border={'bottom': (GOLD, 12, 18)}),
    'Prog Cover Line':   dict(size=10.5, color=CREAM, after=6, line=1.15),
    'Prog Cover Stat':   dict(size=46, bold=True, color=GOLD, before=2, after=16, line=1.0),
    'Prog Cover Name':   dict(size=12, bold=True, color=CREAM, after=1),
    'Prog Cover Site':   dict(size=9, color=CREAM_MUTE, after=0),
    'Prog Body':         dict(size=10.5, after=6, line=1.15),
    'Prog Intro':        dict(size=10.5, color=MUTE, after=6, line=1.15),
    'Prog Note':         dict(size=9.5, color=MUTE, after=8),
    'Prog Table Head':   dict(size=8, bold=True, color=GOLD_TXT, caps=True, track=0.4, after=0),
    'Prog Table Key':    dict(size=8.5, bold=True, color=GOLD_TXT, caps=True, track=0.6, after=0),
    'Prog Table Text':   dict(size=9.5, after=0),
    'Prog Table Strong': dict(size=9.5, bold=True, after=0),
    'Prog Table Muted':  dict(size=8.5, color=MUTE, after=0),
    'Prog Table Id':     dict(size=9, bold=True, color=GOLD_TXT, after=0),
    'Prog Table Number': dict(size=11, bold=True, color=GOLD, after=0),
    'Prog Table Hours':  dict(size=9.5, bold=True, after=0, align='right'),
    'Prog Table Total':  dict(size=11, bold=True, color=GOLD_TXT, after=0, align='right'),
    'Prog Box Tag':      dict(size=8, bold=True, color=GOLD_TXT, caps=True, track=1.2, after=0, keep=True, **_BOX),
    'Prog Box Title':    dict(size=12, bold=True, after=4, keep=True, **_BOX),
    'Prog Box Text':     dict(size=9.5, after=3, keep=True, **_BOX),
    'Prog Bullet':       dict(base='List Bullet', size=10, after=4, line=1.15),
}
CHAR_STYLES = {
    'Prog Label':  dict(size=8.5, bold=True, color=GOLD_TXT, caps=True, track=0.6),
    'Prog Strong': dict(bold=True, color=INK),
    'Prog Cover Label': dict(size=8.5, bold=True, color=GOLD, caps=True, track=0.8),
    'Prog Cover Unit':  dict(size=10, bold=True, color=CREAM_MUTE, caps=True, track=1.5),
}
TABLE_STYLE = 'Prog Table'
ALIGN = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'right': WD_ALIGN_PARAGRAPH.RIGHT}


def rgb(hex_):
    return RGBColor.from_string(hex_)


def _font(font, s):
    font.name = FONT
    if 'size' in s: font.size = Pt(s['size'])
    if 'bold' in s: font.bold = s['bold']
    if 'color' in s: font.color.rgb = rgb(s['color'])
    if s.get('caps'): font.all_caps = True


def _track(el, pts):
    rPr = el.get_or_add_rPr()
    sp = OxmlElement('w:spacing'); sp.set(qn('w:val'), str(int(pts * 20))); rPr.append(sp)


def _borders(pPr, sides):
    bdr = OxmlElement('w:pBdr')
    for side in ('top', 'left', 'bottom', 'right'):
        if side in sides:
            color, sz, gap = sides[side]
            e = OxmlElement(f'w:{side}')
            for k, v in (('val', 'single'), ('sz', sz), ('space', gap), ('color', color)):
                e.set(qn(f'w:{k}'), str(v))
            bdr.append(e)
    pPr.append(bdr)


def _shade(pr, fill):
    s = OxmlElement('w:shd')
    s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), fill)
    pr.append(s)


def right_tab(pf):
    pf.tab_stops.add_tab_stop(Cm(WIDTH), WD_TAB_ALIGNMENT.RIGHT)


def make_styles(doc):
    styles = doc.styles
    styles['Normal'].element.get_or_add_rPr()
    for name, s in STYLES.items():
        st = styles[name] if name in [x.name for x in styles] else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        if name.startswith('Prog'):
            st.base_style = styles[s.get('base', 'Normal')]
            st.quick_style = True
        _font(st.font, s)
        if 'track' in s: _track(st.element, s['track'])
        pf = st.paragraph_format
        pf.space_before, pf.space_after = Pt(s.get('before', 0)), Pt(s.get('after', 6))
        if 'line' in s: pf.line_spacing = s['line']
        if 'indent' in s: pf.left_indent, pf.right_indent = (Pt(x) for x in s['indent'])
        if s.get('keep'): pf.keep_with_next = True
        if s.get('new_page'): pf.page_break_before = True
        if 'align' in s: pf.alignment = ALIGN[s['align']]
        if s.get('tab'): right_tab(pf)
        pPr = st.element.get_or_add_pPr()
        if 'border' in s: _borders(pPr, s['border'])
        if 'fill' in s: _shade(pPr, s['fill'])
    styles['Normal'].element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    for name, s in CHAR_STYLES.items():
        st = styles.add_style(name, WD_STYLE_TYPE.CHARACTER)
        st.quick_style = True
        _font(st.font, s)
        if 'track' in s: _track(st.element, s['track'])
    # table style: hairlines between rows, gold-ruled pale header row
    styles.element.append(parse_xml(f'''
<w:style {nsdecls("w")} w:type="table" w:customStyle="1" w:styleId="ProgTable">
  <w:name w:val="{TABLE_STYLE}"/><w:basedOn w:val="TableNormal"/><w:uiPriority w:val="40"/><w:qFormat/>
  <w:pPr><w:spacing w:after="0"/></w:pPr>
  <w:tblPr>
    <w:tblBorders><w:bottom w:val="single" w:sz="4" w:space="0" w:color="{LINE}"/>
      <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{LINE}"/></w:tblBorders>
    <w:tblCellMar><w:top w:w="70" w:type="dxa"/><w:left w:w="90" w:type="dxa"/>
      <w:bottom w:w="70" w:type="dxa"/><w:right w:w="90" w:type="dxa"/></w:tblCellMar>
  </w:tblPr>
  <w:trPr><w:cantSplit/></w:trPr>
  <w:tblStylePr w:type="firstRow"><w:tblPr/><w:trPr><w:tblHeader/></w:trPr>
    <w:tcPr><w:tcBorders><w:bottom w:val="single" w:sz="8" w:space="0" w:color="{GOLD}"/></w:tcBorders>
      <w:shd w:val="clear" w:color="auto" w:fill="{PALE}"/></w:tcPr></w:tblStylePr>
</w:style>'''))


# ---------- writing helpers (no direct formatting: styles do the look) ----------

def P(doc, style, *parts):
    """Paragraph from parts: plain strings, or (char_style, text) tuples."""
    p = doc.add_paragraph(style=style)
    add(p, *parts)
    return p


FR = False  # French typography, set per programme language


def typo(text):
    """French: narrow no-break space before : ; ? ! and a no-break space between a number and h."""
    if not FR or not isinstance(text, str): return text
    text = re.sub(r' ([:;?!])', '\u202f\\1', text)
    return re.sub(r'(\d) h\b', '\\1\u00a0h', text)


def add(p, *parts):
    for part in parts:
        if isinstance(part, tuple):
            p.add_run(typo(part[1]), style=part[0])
        elif part:
            p.add_run(typo(part))


def page_break(doc):
    p = doc.add_paragraph(style='Prog Body')
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_break(WD_BREAK.PAGE)


def table(doc, widths, header, rows, together=False):
    """rows: lists of (paragraph_style, parts) per cell."""
    t = doc.add_table(rows=0, cols=len(widths))
    t.style = doc.styles[TABLE_STYLE]
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    look = t._tbl.tblPr.find(qn('w:tblLook'))
    if look is None:
        look = OxmlElement('w:tblLook'); t._tbl.tblPr.append(look)
    for k, v in (('firstRow', 1 if header else 0), ('lastRow', 0), ('firstColumn', 0),
                 ('lastColumn', 0), ('noHBand', 1), ('noVBand', 1)):
        look.set(qn(f'w:{k}'), str(v))
    body = ([[('Prog Table Head', [h]) for h in header]] if header else []) + rows
    numeric = [style == 'Prog Table Hours' for style, _ in rows[0]]
    for r, cells in enumerate(body):
        row = t.add_row()
        for c, w, (style, parts), right in zip(row.cells, widths, cells, numeric):
            c.width = Cm(w)
            p = c.paragraphs[0]
            p.style = doc.styles[style]
            if header and r == 0 and right: p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if together: p.paragraph_format.keep_with_next = True
            add(p, *parts)
    for col, w in zip(t.columns, widths):
        col.width = Cm(w)
    return t


def spacer(doc):
    return doc.add_paragraph(style='Prog Body')


DEC = '.'  # decimal separator, set per programme language


def hrs(h):
    return f'{h:g}\u00a0h'.replace('.', DEC) if h is not None else '—'


def logo_png(color=INK):
    """The site logo is white on transparent; recolour it for a white page."""
    im = Image.open(LOGO).convert('RGBA')
    solid = Image.new('RGBA', im.size, '#' + color)
    solid.putalpha(im.getchannel('A'))
    buf = io.BytesIO(); solid.save(buf, 'PNG'); buf.seek(0)
    return buf


# The site's Nabeul tiles (main.js, ART), 100x100 line art
_FR = '<rect x="3" y="3" width="94" height="94" rx="2"/><rect x="8" y="8" width="84" height="84" rx="1"/>'
_CO = '<path d="M23 3A20 20 0 0 1 3 23M77 3A20 20 0 0 0 97 23M3 77A20 20 0 0 1 23 97M97 77A20 20 0 0 1 77 97"/>'
def _sq(r, rot): return f'<rect x="{50-r}" y="{50-r}" width="{2*r}" height="{2*r}" transform="rotate({rot} 50 50)"/>'
def _pet(n, cy, rx, ry, off=0): return ''.join(f'<ellipse cx="50" cy="{cy}" rx="{rx}" ry="{ry}" transform="rotate({off + k*360/n} 50 50)"/>' for k in range(n))
def _dia(x, y, r): return f'<path d="M{x} {y-r}L{x+r} {y}L{x} {y+r}L{x-r} {y}z"/><circle cx="{x}" cy="{y}" r="3"/>'
TILES = [
    _FR + _CO + _sq(27, 0) + _sq(27, 45) + '<circle cx="50" cy="50" r="15"/>' + _sq(8, 0) + _sq(8, 45),
    _FR + _CO + _pet(4, 28, 9, 19) + _pet(4, 32, 5, 12, 45) + '<circle cx="50" cy="50" r="6"/>',
    _FR + ''.join(_dia(x, y, 20) for x, y in [(25, 25), (75, 25), (25, 75), (75, 75), (50, 50)]),
    _FR + _CO + '<circle cx="50" cy="50" r="38"/><circle cx="50" cy="50" r="31"/><circle cx="50" cy="50" r="22"/>' + _sq(14, 0) + _sq(14, 45) + '<circle cx="50" cy="50" r="5"/>',
    _FR + _CO + '<path d="M50 14C70 32 70 68 50 86 30 68 30 32 50 14z"/><path d="M14 50C32 30 68 30 86 50 68 70 32 70 14 50z"/>' + _pet(4, 39, 4, 7, 45) + '<circle cx="50" cy="50" r="7"/><circle cx="50" cy="50" r="2"/>',
]
COVER_ROWS = [0.25, 0.20, 0.15]  # tile rows from the top, fading out


TILE_COLS = 8  # tiles across the page; each row is 210/8 = 26.25 mm high


def tiles_png(dpi=200):
    """The cover's tile band: rows of the site's Nabeul tiles fading out, on transparent."""
    import fitz  # PyMuPDF renders the SVG
    W, s = 210, 210 / TILE_COLS
    H = s * len(COVER_ROWS)
    body = ''
    for r, op in enumerate(COVER_ROWS):
        for c in range(TILE_COLS):
            k = (s - 1.6) / 100
            body += (f'<g transform="translate({c*s + 0.8} {r*s + 0.8}) scale({k})" fill="none" stroke="#{GOLD}" '
                     f'stroke-opacity="{op}" stroke-width="{0.4/k:.3f}" stroke-linejoin="round" stroke-linecap="round">'
                     f'{TILES[(c + 2 * r) % 5]}</g>')
    svg = f'<svg xmlns="http://www.w3.org/2000/svg" width="{W}mm" height="{H:.2f}mm" viewBox="0 0 {W} {H:.2f}">{body}</svg>'
    pix = fitz.open('svg', svg.encode())[0].get_pixmap(dpi=dpi, alpha=True)
    return io.BytesIO(pix.tobytes('png')), Cm(21), Cm(H / 10)


def night_png():
    buf = io.BytesIO(); Image.new('RGB', (8, 8), '#' + NIGHT).save(buf, 'PNG'); buf.seek(0)
    return buf


def behind_text(p, image, width, height, name, locked, x=0, y=0):
    """Picture placed on the page (offsets from its top-left corner), behind the text."""
    inline = p.add_run().add_picture(image, width=width, height=height)._inline
    anchor = parse_xml(
        f'<wp:anchor {nsdecls("wp", "a", "pic", "r")} distT="0" distB="0" distL="0" distR="0" simplePos="0" '
        f'relativeHeight="{2 if not locked else 1}" behindDoc="1" locked="{int(locked)}" layoutInCell="1" allowOverlap="1">'
        f'<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="page"><wp:posOffset>{int(x)}</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="page"><wp:posOffset>{int(y)}</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{inline.extent.cx}" cy="{inline.extent.cy}"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
        f'<wp:wrapNone/><wp:docPr id="{900 + locked}" name="{name}"/><wp:cNvGraphicFramePr/></wp:anchor>')
    anchor.append(inline.graphic)
    inline.getparent().replace(inline, anchor)


def field(p, code):
    r = p.add_run()
    for kind, text in (('begin', None), ('instr', code), ('separate', None), ('t', '1'), ('end', None)):
        if kind in ('instr', 't'):
            e = OxmlElement('w:instrText' if kind == 'instr' else 'w:t'); e.text = text
            e.set(qn('xml:space'), 'preserve')
        else:
            e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), kind)
        r._r.append(e)


# ---------- page furniture ----------

def setup(doc, title, L=LABELS):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2)
    sec.header_distance = sec.footer_distance = Cm(1)
    sec.different_first_page_header_footer = True

    for name, top in (('Header', False), ('Footer', True)):
        st = doc.styles[name]
        _font(st.font, dict(size=8, color=MUTE))
        pPr = st.element.get_or_add_pPr()
        for tabs in pPr.findall(qn('w:tabs')):  # drop the built-in centre/right tabs
            pPr.remove(tabs)
        right_tab(st.paragraph_format)
        _borders(st.element.get_or_add_pPr(),
                 {'top': (LINE, 4, 6)} if top else {'bottom': (GOLD, 6, 4)})

    hp = sec.header.paragraphs[0]
    add(hp, ('Prog Label', L['kicker']), '\t' + title)

    p = sec.footer.paragraphs[0]
    p.add_run().add_picture(logo_png(), height=Cm(0.3))
    add(p, '  ', ('Prog Strong', NAME), '  ·  ' + SITE + '\t' + L['page'] + ' ')
    field(p, 'PAGE'); p.add_run(f" {L['of']} "); field(p, 'NUMPAGES')

    # cover: the black page, locked behind everything (first-page header); no cover footer
    hp = sec.first_page_header.paragraphs[0]
    hp.style = doc.styles['Prog Cover Line']  # not 'Header': no gold rule on the cover
    behind_text(hp, night_png(), Cm(21), Cm(29.7), 'Cover black', locked=True)


AUTHOR_Y = Twips(14592)  # top of the author block, from the top of the page (set by hand in Word)


def author_block(doc):
    """Logo beside name and site; a floating table, draggable in Word by its corner handle."""
    t = doc.add_table(rows=1, cols=2)
    t.autofit = False
    t._tbl.tblPr.append(parse_xml(
        f'<w:tblpPr {nsdecls("w")} w:leftFromText="0" w:rightFromText="0" w:vertAnchor="page" '
        f'w:horzAnchor="margin" w:tblpY="{int(AUTHOR_Y.twips)}"/>'))
    logo_cell, who = t.rows[0].cells
    logo_cell.width, who.width = Cm(1.9), Cm(9)
    who.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = logo_cell.paragraphs[0]; p.style = doc.styles['Prog Cover Name']
    p.add_run().add_picture(logo_png(CREAM), height=Cm(1.35))
    p = who.paragraphs[0]; p.style = doc.styles['Prog Cover Name']; p.add_run(NAME)
    p = who.add_paragraph(style='Prog Cover Site'); p.add_run(SITE)


# ---------- checks ----------

def check(g):
    """Hours rules shared by every programme: chapter sums, 3 h sessions, buffer size, stated duration."""
    errors, ends, run = [], set(), 0
    for ph in g['phases']:
        s = sum(c['hours'] for c in ph['chapters'])
        if ph['hours'] is not None and abs(s - ph['hours']) > 1e-9:
            errors.append(f"phase '{ph['name']}': chapters add up to {s:g} h, phase says {ph['hours']:g} h")
        for c in ph['chapters']:
            run += c['hours']; ends.add(round(run, 2))
    if run % 3:
        errors.append(f'{run:g} h of chapters is not a whole number of 3 h sessions')
    split = [f'{x} h' for x in range(3, int(run), 3) if x not in ends]
    if split:
        errors.append('a 3 h session ends mid-chapter at ' + ', '.join(split) + ' (pair the 1.5 h chapters)')
    buffer = g.get('buffer') or 0
    want = 3 if run <= 24 else 6 if run <= 45 else 9
    if buffer != want:
        errors.append(f'buffer is {buffer:g} h; the rule gives {want} h for {run:g} h of chapters')
    total = run + buffer
    if g.get('hours') not in (None, total):
        errors.append(f"cover hours {g['hours']} differ from the table total {total:g}")
    for f in g.get('facts', []):
        m = re.match(r'(\d+(?:[.,]\d+)?)\s*h\b', f['value'])
        if f['label'] in ('Duration', 'Durée') and (not m or float(m[1].replace(',', '.')) != total):
            errors.append(f"'{f['label']}' fact does not start with the total ({total:g} h)")
    if errors:
        raise SystemExit(f"{g['slug']}:\n  " + '\n  '.join(errors))


# ---------- the document ----------

def build_doc(g):
    global DEC, FR
    fr = FR = g.get('lang') == 'fr'
    DEC = ',' if fr else '.'
    L = {**(LABELS_FR if fr else LABELS), **g.get('labels', {})}
    buffer = g.get('buffer') or 0
    doc = Document()
    make_styles(doc)
    setup(doc, g['title'], L)
    cp = doc.core_properties
    cp.title, cp.author, cp.subject = g['title'], NAME, L['kicker']

    # cover (page 1, on the black background)
    k = P(doc, 'Prog Cover Kicker', L['kicker'])
    img, w, h = tiles_png()
    behind_text(k, img, w, h, 'Cover tiles', locked=False)
    # no line break inside "(RHEL 10)" or before "&"
    title = re.sub(r'\([^)]*\)', lambda m: m[0].replace(' ', ' '), g['title']).replace(' &', ' &')
    P(doc, 'Prog Cover Title', title)
    if g.get('subtitle'): P(doc, 'Prog Cover Sub', g['subtitle'])
    total = g.get('hours') or sum(ph['hours'] or 0 for ph in g['phases']) + buffer
    P(doc, 'Prog Cover Stat', f'{total:g}', ('Prog Cover Unit', '  ' + L['hours']))
    if g.get('audience_line'):
        P(doc, 'Prog Cover Line', ('Prog Cover Label', L['for']), '   ' + g['audience_line'])
    if g.get('certifications'):
        # no-break spaces and hyphens inside each name: lines break only between certifications
        certs = (c.replace(' ', '\u00a0').replace('-', '\u2011') for c in g['certifications'])
        P(doc, 'Prog Cover Line', ('Prog Cover Label', L['certs']), '   ' + '  ·  '.join(certs))
    author_block(doc)
    page_break(doc)

    P(doc, 'Heading 2', L['overview']).paragraph_format.space_before = Pt(0)
    P(doc, 'Prog Body', g['overview'])

    P(doc, 'Heading 2', L['glance'])
    rows = [[('Prog Table Number', [f'{i:02d}']), ('Prog Table Strong', [ph['name']]),
             ('Prog Table Muted', [ph['focus']]), ('Prog Table Hours', [hrs(ph['hours'])])]
            for i, ph in enumerate(g['phases'], 1)]
    if buffer:
        rows.append([('Prog Table Text', []), ('Prog Table Strong', [L['buffer']]),
                     ('Prog Table Muted', [L['buffer_focus']]), ('Prog Table Hours', [hrs(buffer)])])
    total = sum(ph['hours'] or 0 for ph in g['phases']) + buffer
    rows.append([('Prog Table Text', []), ('Prog Table Strong', [L['total']]),
                 ('Prog Table Text', []), ('Prog Table Total', [hrs(total)])])
    t = table(doc, [1.2, 5.2, 8.8, 1.8], L['glance_cols'], rows, together=True)
    for c in t.rows[-1].cells:  # total row: ink rule above
        tcPr = c._tc.get_or_add_tcPr()
        tcPr.append(parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="10" w:space="0" w:color="{INK}"/></w:tcBorders>'))
    spacer(doc)

    if g.get('facts'):
        P(doc, 'Heading 2', L['facts'])
        table(doc, [4.2, WIDTH - 4.2], None,
              [[('Prog Table Key', [f['label']]), ('Prog Table Text', [f['value']])] for f in g['facts']])

    # phases: each starts on a new page through the Prog Kicker style, so no spacer is left
    # before it (an empty paragraph there could spill onto a blank page)
    n = 0
    for i, ph in enumerate(g['phases'], 1):
        P(doc, 'Prog Kicker', f"{L['phase']} {i:02d}\t{hrs(ph['hours'])}")
        P(doc, 'Heading 1', ph['name'])
        if ph.get('intro'): P(doc, 'Prog Intro', ph['intro'])
        if ph.get('note'): P(doc, 'Prog Note', ('Prog Label', L['prereq']), '  ' + ph['note'])
        rows = []
        for c in ph['chapters']:
            content = c['content']
            parts = [('Prog Label', L['maths']), '  ' + content[5:].strip()] if content.startswith('MATH:') else [content]
            rows.append([('Prog Table Id', [c['id']]), ('Prog Table Strong', [c['title']]),
                         ('Prog Table Text', parts), ('Prog Table Muted', [c['format']]),
                         ('Prog Table Hours', [hrs(c['hours'])])])
        table(doc, [1.0, 3.9, 8.6, 2.0, 1.5], L['chapter_cols'], rows)
        pr = ph.get('project')
        if ph.get('math') or pr: spacer(doc)
        if ph.get('math'):
            P(doc, 'Prog Box Tag', L['maths_box'])
            last = P(doc, 'Prog Box Text', ph['math'])
            last.paragraph_format.keep_with_next = False
            if pr: spacer(doc)
        if pr:
            n += 1
            single = sum(1 for q in g['phases'] if q.get('project')) == 1
            P(doc, 'Prog Box Tag', L['final_project'] if single else f"{L['project']} {n}")
            P(doc, 'Prog Box Title', pr['title'])
            last = P(doc, 'Prog Box Text', pr['brief'])
            if pr.get('stack'): last = P(doc, 'Prog Box Text', ('Prog Label', L['stack']), '  ' + pr['stack'])
            if pr.get('deliverable'): last = P(doc, 'Prog Box Text', ('Prog Label', L['deliverable']), '  ' + pr['deliverable'])
            last.paragraph_format.keep_with_next = False

    pf = g.get('portfolio')
    if pf:
        P(doc, 'Prog Kicker', L['portfolio_kicker'])
        P(doc, 'Heading 1', L['portfolio'])
        P(doc, 'Prog Intro', pf['intro'])
        rows = [[('Prog Table Number', [f'{i:02d}']), ('Prog Table Strong', [r['name']]),
                 ('Prog Table Text', [r['shows']]), ('Prog Table Muted', [r['phase']]),
                 ('Prog Table Muted', [r['build']])] for i, r in enumerate(pf['rows'], 1)]
        table(doc, [1.2, 3.8, 8.0, 2.0, 2.0], L['portfolio_cols'], rows)
        spacer(doc)
    if g.get('notes'):
        P(doc, 'Heading 2', L['notes'])
        for note in g['notes']:
            P(doc, 'Prog Bullet', note)
    if doc.element.body[-2].tag == qn('w:tbl'):  # Word needs a paragraph after a closing table
        spacer(doc)
    return doc


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load_built():
    return json.loads(BUILT.read_text()) if BUILT.exists() else {}


def build(path, force=False, out=None):
    g = json.loads(path.read_text(encoding='utf-8'))
    check(g)
    target = out or OUT / g['file']
    built = load_built()
    if out is None and target.exists() and not force and built.get(g['file']) != sha(target):
        raise SystemExit(f'{target.name} was edited in Word since the last build. '
                         f'Run scripts/sync_programmes.py first (or --force to discard the edits).')
    target.parent.mkdir(parents=True, exist_ok=True)
    build_doc(g).save(target)
    if out is None:
        built[g['file']] = sha(target)
        BUILT.write_text(json.dumps(built, indent=1) + '\n')
    return target


if __name__ == '__main__':
    args = sys.argv[1:]
    force = '--force' in args
    names = [a for a in args if not a.startswith('--')]
    files = [SRC / f'{n}.json' for n in names] if names else sorted(f for f in SRC.glob('*.json') if not f.name.startswith(('.', '_')))
    for f in files:
        print('built', build(f, force).relative_to(ROOT))
